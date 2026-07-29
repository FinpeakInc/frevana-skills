#!/usr/bin/env python3

import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path


SKILL_DIR = Path(__file__).resolve().parents[1]
PY_SCRIPT = SKILL_DIR / "scripts" / "local_knowledge.py"
PARSER_MODULES = ("pypdf", "docx", "openpyxl")


CHROMADB_STUB = """
_collections = {}


class Collection:
    def __init__(self):
        self.rows = {}

    def upsert(self, ids, documents, metadatas, embeddings):
        for index, row_id in enumerate(ids):
            self.rows[row_id] = (
                documents[index],
                metadatas[index],
                embeddings[index],
            )

    def count(self):
        return len(self.rows)


class PersistentClient:
    def __init__(self, path):
        self.path = path

    def delete_collection(self, name):
        key = (self.path, name)
        if key not in _collections:
            raise KeyError(name)
        del _collections[key]

    def get_or_create_collection(self, name, metadata=None):
        return _collections.setdefault((self.path, name), Collection())
"""


SENTENCE_TRANSFORMERS_STUB = """
class SentenceTransformer:
    def __init__(self, *args, **kwargs):
        pass

    def encode(self, texts, **kwargs):
        return [[float(len(text)), 1.0] for text in texts]
"""


def parser_dependencies_available() -> bool:
    return all(importlib.util.find_spec(module) is not None for module in PARSER_MODULES)


@unittest.skipUnless(
    parser_dependencies_available(),
    "real parser integration test requires pypdf, python-docx, and openpyxl",
)
class RealDocumentParserIntegrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="local-knowledge-real-docs-")
        self.root = Path(self.temp_dir.name)
        self.source_dir = self.root / "source"
        self.data_dir = self.root / "data"
        self.output_path = self.root / "result.json"
        self.stub_dir = self.root / "stubs"
        self.source_dir.mkdir()
        self.stub_dir.mkdir()
        self._write(self.stub_dir / "chromadb.py", CHROMADB_STUB)
        self._write(
            self.stub_dir / "sentence_transformers.py",
            SENTENCE_TRANSFORMERS_STUB,
        )
        self._create_pdf(self.source_dir / "guide.pdf", "REAL_PDF_OK")
        self._create_docx(self.source_dir / "notes.docx", "REAL_DOCX_OK")
        self._create_xlsx(self.source_dir / "metrics.xlsx", "REAL_XLSX_OK")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    @staticmethod
    def _write(path: Path, content: str) -> None:
        path.write_text(textwrap.dedent(content).lstrip(), encoding="utf-8")

    @staticmethod
    def _create_pdf(path: Path, text: str) -> None:
        from pypdf import PdfWriter
        from pypdf.generic import (
            DecodedStreamObject,
            DictionaryObject,
            NameObject,
        )

        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_ref}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(
            f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
        with path.open("wb") as handle:
            writer.write(handle)

    @staticmethod
    def _create_docx(path: Path, text: str) -> None:
        import docx

        document = docx.Document()
        document.add_paragraph(text)
        document.save(path)

    @staticmethod
    def _create_xlsx(path: Path, text: str) -> None:
        import openpyxl

        workbook = openpyxl.Workbook()
        workbook.active.append(("name", "value"))
        workbook.active.append((text, 3))
        workbook.save(path)

    def test_real_pdf_docx_xlsx_files_are_indexed(self) -> None:
        env = os.environ.copy()
        existing = env.get("PYTHONPATH")
        paths = [str(self.stub_dir)]
        if existing:
            paths.append(existing)
        env["PYTHONPATH"] = os.pathsep.join(paths)

        result = subprocess.run(
            [
                sys.executable,
                str(PY_SCRIPT),
                "index",
                "--data-dir",
                str(self.data_dir),
                "--path",
                str(self.source_dir),
                "--output",
                str(self.output_path),
            ],
            cwd=self.root,
            env=env,
            text=True,
            capture_output=True,
            check=False,
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = json.loads(self.output_path.read_text(encoding="utf-8"))
        self.assertEqual(payload["stats"]["file_count"], 3)
        self.assertEqual(payload["stats"]["skipped_count"], 0)
        chunks_path = Path(payload["index_dir"]) / "chunks.jsonl"
        chunks = chunks_path.read_text(encoding="utf-8")
        self.assertIn("REAL_PDF_OK", chunks)
        self.assertIn("REAL_DOCX_OK", chunks)
        self.assertIn("REAL_XLSX_OK", chunks)


if __name__ == "__main__":
    unittest.main(verbosity=2)
